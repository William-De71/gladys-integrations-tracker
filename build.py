#!/usr/bin/env python3
"""Build the tracker: Pierre-Gilles' Home Assistant list ↔ Gladys external integrations.

Everything is found again on every run, with nothing to maintain by hand:
  - Pierre-Gilles' gist (3_full-list.csv), and data/ha-forum-juillet.csv: the popular HA integrations of his
    July analysis (forum post #1) that the gist left out because Gladys already covered them;
  - the live Gladys store index (index.json);
  - GitHub repositories with the gladys-assistant-integration topic (published, or about to be);
  - GitHub repositories named "gladys…" that contain a gladys-assistant-integration.json manifest
    (every repository created from the integration template has one), or that have the
    gladys-integration-wip topic: integrations in progress, which the store ignores;
  - archived, deprecated (description starting with "Deprecated") or deleted repositories, checked on GitHub itself
    (the store index can lag behind);
  - claims posted in the forum topic ("je prends X"), until a repository shows up;
  - HA list ↔ Gladys matches: integrations whose names are the same are matched automatically.

Optional corrections, never required:
  - data/mapping.txt: HA list number | repository | E (equivalent), P (partial) or - (wrong automatic match);
  - data/claims.txt: claims the forum reading misses, or ideas freed again ("-" as the claimant);
  - data/notes.txt: warnings shown under an idea (terms of use...), in French and English;
  - data/extra-repos.txt: Gladys repositories the searches cannot find.

Output, in site/: data.json (read by index.html), one CSV per table, and an .xlsx workbook.
"""
import csv
import html
import io
import json
import shutil
import subprocess
import re
import sys
import unicodedata
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from functools import cache
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
OUT = ROOT / "site"

GIST_ID = "9c4ba77a0faa47c4ea67c74351db9094"
GIST_URL = f"https://gist.github.com/Pierre-Gilles/{GIST_ID}"
# The raw URL needs no token: the Actions GITHUB_TOKEN cannot read the gists API.
GIST_CSV = f"https://gist.githubusercontent.com/Pierre-Gilles/{GIST_ID}/raw/3_full-list.csv"
STORE_INDEX = "https://integration-store-storage.gladysassistant.com/index.json"
STORE_PAGE = "https://gladysassistant.com/fr/docs/integrations/external/"
FORUM = "https://community.gladysassistant.com"
FORUM_TOPIC = f"{FORUM}/t/10385"
# How a developer claims an integration in the forum topic, and what ends the claim within the same sentence.
# "en cours" alone is not one: "appli en cours", "téléchargements en cours"…
CLAIM_WORDS = (r"(je prends|je reprends|j ai lance le dev|je lance|encore en cours|en cours de dev|je m occupe de"
               r"|je pars sur|je me lance sur)")
CLAIM_STOPS = (r"(reste chez|restent chez|pris par|prise par|je n y touche pas|deja pris|deja fait|deja dans le store"
               r"|n est pas|sortie|sorties|est sortie)")
# HA list names that also come up as everyday words in claims ("vu sous une vidéo YouTube", "vérifié sur GitHub").
CLAIM_IGNORED = {"youtube", "github"}
# Without numbers, names are only read this far after "je prends": long posts go on with descriptions.
CLAIM_WINDOW = 500
# Status posts list ideas without claiming them.
CLAIM_AUTHORS_IGNORED = {"pierre-gilles"}
TOPIC = "gladys-assistant-integration"
# Claims an integration without the store topic: the store would list an empty repository with it.
WIP_TOPIC = "gladys-integration-wip"
MANIFEST = "gladys-assistant-integration.json"
# External integrations started in summer 2026: older "gladys" repositories are Gladys 4 modules.
NAME_SEARCH = "gladys in:name fork:false created:>2026-06-01"
USER_AGENT = "gladys-integrations-tracker"
# Where to start an update by hand: Actions tab → "Run workflow" (needs write access to the repository).
WORKFLOW_PAGE = "https://github.com/William-De71/gladys-integrations-tracker/actions/workflows/build.yml"

# Status keys, from the most to the least advanced. The page and the workbook translate them.
STORE, TOPIC_ONLY, NO_TOPIC, INACTIVE, ARCHIVED = "store", "topic", "no_topic", "inactive", "archived"
NATIVE, CLAIMED, FREE = "native", "claimed", "free"
STATUS_FR = {
    STORE: "Publiée (store)",
    TOPIC_ONLY: "Topic store posé, pas encore indexée",
    NO_TOPIC: "En cours (dépôt GitHub)",
    INACTIVE: "Hors store, inactive",
    CLAIMED: "Annoncée sur le forum",
    FREE: "Libre",
    NATIVE: "Native Gladys",
    ARCHIVED: "Archivée, dépréciée ou supprimée",
}
RANK = [STORE, TOPIC_ONLY, NO_TOPIC, INACTIVE, ARCHIVED]
# A repository outside the store with no commit for this long is no longer "in progress".
INACTIVE_AFTER = timedelta(days=30)
QUALITY_FR = {"E": "Équivalente", "P": "Partielle", "A": "Automatique (même nom)"}


def fetch(url):
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=30) as r:
        return r.read().decode("utf-8")


def gh_api(*args):
    result = subprocess.run(["gh", "api", *args], capture_output=True, text=True)
    if result.returncode != 0:
        raise RuntimeError(f"gh api {' '.join(args)}: {result.stderr.strip()}")
    return result.stdout


def read_table(path):
    rows = []
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#"):
            rows.append([c.strip() for c in line.split("|")])
    return rows


def load_ha_list():
    """Pierre-Gilles' list, then the July analysis entries (ids F1, F2…), in the same columns."""
    ideas = list(csv.DictReader(io.StringIO(fetch(GIST_CSV))))
    for h in ideas:
        h["_source"], h["_native"], h["_note"] = "list", False, ""
    with open(DATA / "ha-forum-juillet.csv", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            ideas.append({
                "#": r["id"], "Priority": r["priority"], "Integration": r["name"], "Category": r["category"],
                "HA usage (% of installations)": r["ha_usage_pct"], "Access": r["access"],
                "No hardware needed": "", "Area": "Worldwide", "Status": "Available",
                "Description": r["description"], "Reference link": r["link"],
                "_source": "forum", "_native": r["native_gladys"] == "oui", "_note": r["forum_table"],
            })
    return ideas


def order(num):
    """Sort key: the July analysis entries (F1, F2…) first, then the list from 1 to 1190."""
    return int(num) if num.isdigit() else int(num[1:]) - 100000


def exists(url):
    try:
        req = urllib.request.Request(url, method="HEAD", headers={"User-Agent": USER_AGENT})
        with urllib.request.urlopen(req, timeout=10) as r:
            return r.status == 200
    except Exception:
        return False


def read_repo(repo, branch):
    """Manifest and cover of a repository that is not in the store index yet."""
    raw = f"https://raw.githubusercontent.com/{repo}/{branch}/"
    try:
        manifest = json.loads(fetch(raw + MANIFEST))
    except Exception:
        manifest = None
    cover = next((raw + n for n in ("cover.png", "cover.jpg") if exists(raw + n)), "")
    if manifest and not cover and str(manifest.get("cover_image", "")).startswith("https://"):
        cover = manifest["cover_image"]
    return manifest, cover


def readable(repo_name):
    """gladys-truenas → TrueNAS is out of reach, but at least Truenas."""
    stem = re.sub(r"(?i)^gladys[-_]?|[-_]?gladys$|[-_]integration$|^integration[-_]", "", repo_name) or repo_name
    return " ".join(w[:1].upper() + w[1:] for w in re.split(r"[-_]+", stem))


def search_repos(query):
    out = gh_api("-X", "GET", "search/repositories", "-f", f"q={query}", "-f", "per_page=100",
                 "--paginate", "--jq", ".items[]")
    return [json.loads(line) for line in out.splitlines() if line.strip()]


def load_gladys():
    index = json.loads(fetch(STORE_INDEX))
    repos = {}
    for item in index["integrations"]:
        m = item["manifest"]
        repo = item["repo_url"].replace("https://github.com/", "")
        desc = m.get("description") or {}
        repos[repo.lower()] = dict(
            repo=repo, name=m["name"], author=repo.split("/")[0],
            desc_fr=desc.get("fr") or desc.get("en", ""), desc_en=desc.get("en") or desc.get("fr", ""),
            type=m.get("type", ""), version=m.get("version", ""), cover=item.get("cover_url", ""),
            status=STORE, since=(item.get("first_seen_at") or "")[:10],
            pushed=(item.get("github") or {}).get("pushed_at", "")[:10],
        )

    # Listed whatever their content: repositories with the store topic or the work-in-progress topic.
    # Found by name: listed only when they hold a Gladys manifest, to leave out unrelated "gladys" repositories.
    candidates, wip = {}, set()
    for r in search_repos(f"topic:{TOPIC}"):
        candidates.setdefault(r["full_name"].lower(), (r, TOPIC_ONLY))
    for r in search_repos(f"topic:{WIP_TOPIC}"):
        candidates.setdefault(r["full_name"].lower(), (r, NO_TOPIC))
        wip.add(r["full_name"].lower())
    for r in search_repos(NAME_SEARCH):
        candidates.setdefault(r["full_name"].lower(), (r, NO_TOPIC))
    for line in (DATA / "extra-repos.txt").read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#") and line.lower() not in candidates:
            try:
                candidates[line.lower()] = (json.loads(gh_api(f"repos/{line}")), NO_TOPIC)
            except RuntimeError:
                print(f"repository not found: {line}", file=sys.stderr)
    todo = [(r, status) for key, (r, status) in candidates.items() if key not in repos]

    def inspect(item):
        r, status = item
        return r, status, read_repo(r["full_name"], r.get("default_branch", "main"))

    with ThreadPoolExecutor(16) as ex:
        for r, status, (manifest, cover) in ex.map(inspect, todo):
            if status == NO_TOPIC and manifest is None and r["full_name"].lower() not in wip:
                continue
            m = manifest or {}
            # Created from the integration template and not filled in yet: only the repository name means something.
            template = "template" in str(m.get("name", "")).lower()
            if template:
                m = {}
            if status == NO_TOPIC and datetime.now(timezone.utc) - parse_date(r["pushed_at"]) > INACTIVE_AFTER:
                status = INACTIVE
            desc = m.get("description") if isinstance(m.get("description"), dict) else {}
            fallback = r.get("description") or ""
            repos[r["full_name"].lower()] = dict(
                repo=r["full_name"], name=m.get("name") or readable(r["name"]), author=r["owner"]["login"],
                desc_fr=desc.get("fr") or desc.get("en") or fallback,
                desc_en=desc.get("en") or desc.get("fr") or fallback,
                type=m.get("type", ""), version=m.get("version", ""), cover=cover, status=status,
                since=r["created_at"][:10], pushed=r["pushed_at"][:10], template=template,
            )
    for key in archived_or_gone([g["repo"] for g in repos.values()]):
        repos[key]["status"] = ARCHIVED
    return repos, index.get("generated_at", "")


def parse_date(iso):
    return datetime.fromisoformat(iso.replace("Z", "+00:00"))


def covers(g):
    """Whether an integration still covers its idea: an archived repository no longer does, and neither does an
    untouched copy of the template once inactive, so that the idea is free again."""
    return g["status"] != ARCHIVED and not (g["status"] == INACTIVE and g.get("template"))


# A repository its author deprecated without archiving it says so at the start of its description.
DEPRECATED = re.compile(r"\s*(deprecated|obsol[eè]te|d[ée]pr[ée]ci[ée]e?)\b", re.I)


def archived_or_gone(names):
    """Repositories archived, deprecated or deleted on GitHub, as lowercase owner/name keys."""
    found = set()
    for start in range(0, len(names), 50):
        chunk = names[start:start + 50]
        fields = " ".join(
            f'r{i}: repository(owner: {json.dumps(n.split("/")[0])}, name: {json.dumps(n.split("/")[1])}) {{ isArchived description }}'
            for i, n in enumerate(chunk))
        # A deleted repository makes GraphQL answer with an error next to the data: read the data anyway.
        result = subprocess.run(["gh", "api", "graphql", "-f", f"query={{ {fields} }}"],
                                capture_output=True, text=True)
        try:
            data = json.loads(result.stdout).get("data") or {}
        except json.JSONDecodeError:
            raise RuntimeError(f"gh api graphql: {result.stderr.strip()}")
        for i, n in enumerate(chunk):
            repo = data.get(f"r{i}")
            if repo is None or repo["isArchived"] or DEPRECATED.match(repo["description"] or ""):
                found.add(n.lower())
    return found


def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    return " ".join(t for t in re.sub(r"[^a-z0-9]+", " ", s).split() if t not in NAME_STOPWORDS)


NAME_STOPWORDS = {"gladys", "integration", "integrations", "assistant", "external", "the", "and", "for",
                  "de", "la", "le", "les", "et", "du", "des", "home", "ha"}


def without_domain(name):
    """Healthchecks.io → Healthchecks. Only when 6 letters or more are left: "Alarm.com" or "Fast.com" would leave
    words too common to tell an integration apart."""
    bare = re.sub(r"(?i)\.(io|com|net|org|fr|de|eu|app|dev|co)\b", "", name)
    return bare if bare != name and len(norm(bare).replace(" ", "")) >= 6 else name


def ha_names(h):
    """Names an HA list entry goes by: as written, without its parentheses, and without a domain name."""
    name = h["Integration"]
    bare = re.sub(r"\(.*?\)", "", name)
    return {n for n in (norm(name), norm(bare), norm(without_domain(bare))) if len(n) >= 3}


def ha_parts(h):
    """Each part of a "A / B" HA name: "La Poste / Colissimo / Chronopost tracking" also goes by "La Poste"."""
    name = re.sub(r"\(.*?\)", "", h["Integration"])
    return {n for n in (norm(p) for p in name.split(" / ")) if len(n) >= 3} if " / " in name else set()


def gladys_names(g):
    """Names a Gladys integration goes by: its name, each part of a "A / B" name, and its repository name. Later
    parts also go by the brand of the first one: "IKEA DIRIGERA & TRÅDFRI" covers "IKEA TRÅDFRI"."""
    name = re.sub(r"\(.*?\)", "", g["name"])
    names = {norm(g["name"]), norm(name), norm(without_domain(name))}
    parts = [norm(p) for p in re.split(r" / |, | & ", name)]
    names.update(parts)
    brand = parts[0].split()[0] if len(parts) > 1 and parts[0] else ""
    names.update(f"{brand} {p}" for p in parts[1:] if p and not p.startswith(brand + " "))
    stem = re.sub(r"(?i)^gladys[-_]?|[-_]?gladys$|[-_]integration$|^integration[-_]", "", g["repo"].split("/")[1])
    names.add(norm(stem.replace("-", " ").replace("_", " ")))
    return {n for n in names if len(n) >= 3}


def repo_stem(repo):
    """gladys-bosch-shc → bosch_shc, written like a Home Assistant integration id."""
    stem = re.sub(r"(?i)^gladys[-_]?|[-_]?gladys$|[-_]integration$|^integration[-_]", "", repo.split("/")[1])
    return stem.lower().replace("-", "_")


def ha_id(h):
    """Home Assistant integration id, from the documentation link: .../integrations/bosch_shc/ → bosch_shc."""
    return h["Reference link"].rstrip("/").split("/")[-1].lower().replace("-", "_")


def auto_matches(ha, gladys):
    """Same name on both sides (spaces aside), or a repository named after the Home Assistant integration id (gladys-bosch-shc for
    bosch_shc). Strict on purpose: a wrong match is worse than a missing one. Ids of a single word are left out,
    since they can mislead: Home Assistant's "fritzbox" is FRITZ!SmartHome, not the FRITZ!Box."""
    by_name, by_id = {}, {}
    for h in ha:
        for n in ha_names(h):
            by_name.setdefault(n, set()).add(h["#"])
        if "_" in ha_id(h):
            by_id.setdefault(ha_id(h), set()).add(h["#"])
    # A part of a "A / B" name counts only when it points to that idea alone: "Verisure" in
    # "Securitas Direct / Verisure" is an idea of its own, and "Fire TV" comes up in two ideas.
    parts = {}
    for h in ha:
        for n in ha_parts(h) - ha_names(h):
            parts.setdefault(n, set()).add(h["#"])
    for n, nums in parts.items():
        if n not in by_name and len(nums) == 1:
            by_name[n] = nums
    # Spaces do not tell names apart: "rtl_433" reads as "rtl 433", and gladys-rtl433 as "rtl433".
    compact = {}
    for n, nums in by_name.items():
        compact.setdefault(n.replace(" ", ""), set()).update(nums)
    found = []
    for g in gladys.values():
        nums = set(by_id.get(repo_stem(g["repo"]), set()))
        for n in gladys_names(g):
            nums |= compact.get(n.replace(" ", ""), set())
        found.extend((num, g) for num in nums)
    return found


@cache
def forum_posts():
    topic = json.loads(fetch(FORUM_TOPIC + ".json"))
    ids = topic["post_stream"]["stream"]
    posts = []
    for start in range(0, len(ids), 20):
        query = "&".join(f"post_ids[]={i}" for i in ids[start:start + 20])
        posts += json.loads(fetch(f"{FORUM_TOPIC}/posts.json?{query}"))["post_stream"]["posts"]
    return posts


def spaced(text):
    """Lowercase ASCII words, padded with spaces so that " name " only matches whole words."""
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().lower()
    return " " + re.sub(r"[^a-z0-9]+", " ", text).strip() + " "


def forum_claims(ha):
    """HA list number → the latest forum post claiming it, read in the text after "je prends", up to the next claim
    or to a phrase saying the idea belongs to someone else:
      - numbers of the list written in the post ("Victron Energy GX (P1 n° 74)") are the most reliable: when a claim
        has some, only they count, since its descriptions often name other ideas ("ça couvre Tasmota, BTHome…");
      - otherwise HA names written in full, the longest one winning when they overlap ("Octopus Energy France"
        rather than "Octopus Energy"). A shortened name ("Solcast" for "Solcast PV Forecast") is not read: first
        words mislead too often ("Govee" for Govee Bluetooth, "Transport" for Transport for NSW); data/claims.txt
        covers those."""
    ids = {h["#"] for h in ha}
    names = []
    for h in ha:
        full = {spaced(h["Integration"]), spaced(re.sub(r"\(.*?\)", "", h["Integration"]))}
        names += [(name, h["#"]) for name in full if len(name.strip()) >= 4 and name.strip() not in CLAIM_IGNORED]
    claims = {}
    try:
        posts = forum_posts()
    except Exception as e:
        print(f"forum not read: {e}", file=sys.stderr)
        return claims
    for p in posts:
        if p["username"] in CLAIM_AUTHORS_IGNORED:
            continue
        cooked = re.sub(r'<aside class="quote.*?</aside>', " ", p["cooked"], flags=re.S)
        text = spaced(html.unescape(re.sub(r"<[^>]+>", " ", cooked)))
        for m in re.finditer(CLAIM_WORDS, text):
            rest = text[m.end():]
            end = re.search(CLAIM_WORDS + "|" + CLAIM_STOPS, rest)
            part = " " + (rest[:end.start()] if end else rest) + " "
            # "n° 74", "n°74" ("°" is dropped, not spaced), "nº 74", and the July analysis entries: "n° F39".
            # "au n° 122", "du n° 241" point to a forum post, not to the list.
            numbers = [n.upper() for n in re.findall(r"(?<! au)(?<! du) no? ?(f?\d{1,4}) ", part) if n.upper() in ids]
            if numbers:
                found = set(numbers)
            else:
                part = part[:CLAIM_WINDOW]
                hits = [(name, num) for name, num in names if name in part]
                found = {num for name, num in hits
                         if not any(name != other and name in other for other, _ in hits)}
            for num in found:
                claims[num] = dict(who=p["username"], date=p["created_at"][:10], post=p["post_number"],
                                   url=f"{FORUM_TOPIC}/{p['post_number']}")
    return claims


def build():
    ha = load_ha_list()
    gladys, index_date = load_gladys()

    mapping, rejected = {}, set()
    for num, repo, quality in read_table(DATA / "mapping.txt"):
        g = gladys.get(repo.lower())
        if g is None:
            print(f"mapping skipped, unknown repository: {repo}", file=sys.stderr)
        elif quality == "-":
            rejected.add((num, g["repo"]))
        else:
            mapping.setdefault(num, []).append((g, quality))
    for num, g in auto_matches(ha, gladys):
        if (num, g["repo"]) not in rejected and all(x is not g for x, _ in mapping.get(num, [])):
            mapping.setdefault(num, []).append((g, "A"))
    claims = {num: (who, src) for num, who, src in read_table(DATA / "claims.txt")}
    # "-" frees an idea whose claim was given up: forum posts and the gist status no longer count for it.
    freed = {num for num, (who, _) in claims.items() if who == "-"}
    # Warnings shown under an idea (terms of use, access restrictions), in French and English.
    notes = {num: dict(fr=fr, en=en) for num, fr, en in read_table(DATA / "notes.txt")}
    announced = forum_claims(ha)
    try:
        post_dates = {p["post_number"]: p["created_at"][:10] for p in forum_posts()}
    except Exception:
        post_dates = {}
    # Repositories that no longer cover an idea keep their match for the record, but count neither as coverage
    # nor as duplicates.
    live = {num: [(g, q) for g, q in lst if covers(g)] for num, lst in mapping.items()}

    ha_rows = []
    for h in ha:
        num = h["#"]
        matches = sorted(live.get(num, []), key=lambda m: (RANK.index(m[0]["status"]), "EPA".index(m[1])))
        row = dict(
            n=int(num) if num.isdigit() else num, order=order(num), source=h["_source"], note=h["_note"],
            prio=h["Priority"], name=h["Integration"], category=h["Category"],
            usage=h["HA usage (% of installations)"], access=h["Access"],
            no_hardware=h["No hardware needed"] == "yes", area=h["Area"], link=h["Reference link"],
            desc=h["Description"], status=FREE, quality="", claim="", claim_url="",
            matches=[dict(repo=g["repo"], quality=q) for g, q in matches],
            # Several full matches for one idea; a partial match only complements the others.
            dup=sum(1 for _, q in matches if q != "P") > 1,
            warning=notes.get(num),
        )
        if matches:
            row["status"] = matches[0][0]["status"]
            row["quality"] = matches[0][1]
        elif h["_native"]:
            row["status"] = NATIVE
        elif num in freed:
            pass
        elif num in claims:
            row["status"] = CLAIMED
            who, src = claims[num]
            post = re.search(r"forum #(\d+)", src)
            # Dated like the claims read in the forum when the post is known.
            date = post_dates.get(int(post.group(1))) if post else None
            row["claim"] = f"{who}, {src} ({date})" if date else f"{who} — {src}"
            row["claim_url"] = f"{FORUM_TOPIC}/{post.group(1)}" if post else ""
        elif num in announced and datetime.now(timezone.utc) - parse_date(announced[num]["date"] + "T00:00:00Z") \
                <= INACTIVE_AFTER:
            # A claim without a repository after 30 days frees the idea again, like an inactive repository.
            a = announced[num]
            row["status"] = CLAIMED
            row["claim"] = f"{a['who']}, forum #{a['post']} ({a['date']})"
            row["claim_url"] = a["url"]
        elif h["Status"] != "Available":
            row["status"] = CLAIMED
            row["claim"] = f"gist: {h['Status']}"
        ha_rows.append(row)
    ha_rows.sort(key=lambda r: r["order"])

    # Two Gladys integrations are duplicates when they match the same HA idea, or share a name.
    by_name = {}
    for g in gladys.values():
        g["ha"] = [int(n) if n.isdigit() else n
                   for n in sorted((n for n, lst in mapping.items() if any(x is g for x, _ in lst)), key=order)]
        g["dups"] = set()
        full_name = re.sub(r"[^a-z0-9]+", " ", unicodedata.normalize("NFKD", re.sub(r"\(.*?\)", "", g["name"]))
                           .encode("ascii", "ignore").decode().lower()).strip()
        if covers(g):
            by_name.setdefault(full_name, []).append(g)
    # A partial match complements an integration rather than duplicating it.
    groups = [[g for g, q in lst if q != "P"] for lst in live.values()] + list(by_name.values())
    for group in groups:
        for g in group:
            g["dups"].update(x["repo"] for x in group if x is not g)
    for g in gladys.values():
        g["dups"] = sorted(g["dups"])
    gladys_rows = sorted(gladys.values(), key=lambda g: (RANK.index(g["status"]), g["name"].lower()))

    return dict(
        generated_at=datetime.now(timezone.utc).isoformat(timespec="seconds"),
        store_index_at=index_date,
        sources=dict(ha_list=GIST_URL, store=STORE_PAGE, store_index=STORE_INDEX, forum=FORUM_TOPIC,
                     topic=f"https://github.com/topics/{TOPIC}", workflow=WORKFLOW_PAGE),
        ha=ha_rows,
        gladys=gladys_rows,
    )


def flat_tables(data):
    """Plain values, so that the CSVs work with Google Sheets' IMPORTDATA."""
    by_repo = {g["repo"]: g for g in data["gladys"]}
    head_ha = ["N°", "Prio", "Intégration Home Assistant", "Catégorie", "Usage HA", "Accès",
               "Sans matériel", "Zone", "Doc HA / HACS", "Statut Gladys", "Correspondance",
               "Intégration Gladys", "Auteur", "Dépôt GitHub", "Cover", "Version", "Doublon",
               "Autres dépôts / note", "Description HA", "Source"]
    rows_ha = []
    for h in data["ha"]:
        first = by_repo[h["matches"][0]["repo"]] if h["matches"] else None
        others = [m["repo"] for m in h["matches"][1:]]
        rows_ha.append([
            h["n"], h["prio"], h["name"], h["category"], h["usage"], h["access"],
            "oui" if h["no_hardware"] else "", h["area"], h["link"], STATUS_FR[h["status"]],
            QUALITY_FR.get(h["quality"], ""),
            first["name"] if first else "", first["author"] if first else "",
            f"https://github.com/{first['repo']}" if first else "",
            first["cover"] if first else "", first["version"] if first else "",
            f"oui ({len(h['matches'])})" if h["dup"] else "",
            "; ".join(x for x in (", ".join(others) or h["claim"], (h.get("warning") or {}).get("fr", "")) if x),
            h["desc"],
            "Liste de Pierre-Gilles" if h["source"] == "list" else f"Analyse forum de juillet ({h['note']})",
        ])
    head_g = ["Intégration Gladys", "Auteur", "Dépôt GitHub", "Cover", "Statut", "Type", "Version",
              "Vue depuis", "Dernier push", "Équivalent HA (liste)", "Doublon avec", "Description"]
    names = {h["n"]: h["name"] for h in data["ha"]}
    rows_g = [[g["name"], g["author"], f"https://github.com/{g['repo']}", g["cover"], STATUS_FR[g["status"]],
               g["type"], g["version"], g["since"], g["pushed"],
               ", ".join(f"{n} {names[n]}" for n in g["ha"]) or "hors liste", ", ".join(g["dups"]),
               g["desc_fr"]]
              for g in data["gladys"]]
    return (head_ha, rows_ha), (head_g, rows_g)


def write_xlsx(data, tables, path):
    (head_ha, rows_ha), (head_g, rows_g) = tables
    ideas = [h for h in data["ha"] if h["source"] == "list"]
    july = [h for h in data["ha"] if h["source"] == "forum"]
    n_taken = sum(1 for h in ideas if h["status"] != FREE)
    p1 = [h for h in ideas if h["prio"] == "P1"]
    n_dup = sum(1 for g in data["gladys"] if g["dups"])
    n_store = sum(1 for g in data["gladys"] if g["status"] == STORE)

    def link(url, label=None):
        return f'=HYPERLINK("{url}","{label or url}")' if url else ""

    wb = Workbook()
    ws = wb.active
    ws.title = "Lisez-moi"
    for r in [
        ["Suivi des intégrations externes Gladys"],
        [f"Généré le {data['generated_at'][:16].replace('T', ' ')} UTC "
         f"(index du store : {data['store_index_at'][:16].replace('T', ' ')} UTC)"],
        [],
        ["Liste Home Assistant (gist de Pierre-Gilles)",
         f"{len(ideas)} idées, dont {n_taken} prises ou faites ; "
         f"P1 : {sum(1 for h in p1 if h['status'] != FREE)} / {len(p1)}"],
        ["Analyse forum de juillet (F1, F2…)",
         f"{len(july)} intégrations populaires que la liste a retirées car Gladys les couvrait déjà ; "
         f"{sum(1 for h in july if h['status'] != FREE)} couvertes"],
        ["Intégrations Gladys repérées", f"{len(data['gladys'])} dépôts, dont {n_store} publiés au store"],
        ["Intégrations Gladys en doublon", str(n_dup)],
        [],
        ["Statuts"],
        [STATUS_FR[STORE], "indexée par le store Gladys"],
        [STATUS_FR[NATIVE], "couverte par le cœur de Gladys, sans intégration externe"],
        [STATUS_FR[TOPIC_ONLY], f"dépôt avec le topic {TOPIC}, pas encore indexé par le store (ou refusé par l'indexeur)"],
        [STATUS_FR[NO_TOPIC], f"dépôt créé depuis le template, ou avec le topic {WIP_TOPIC} : en développement"],
        [STATUS_FR[INACTIVE], "dépôt hors store sans commit depuis plus de 30 jours ; une copie du template jamais "
                              "remplie ne compte plus pour l'idée"],
        [STATUS_FR[ARCHIVED], "dépôt archivé, déprécié ou supprimé : ne compte plus pour l'idée"],
        [STATUS_FR[CLAIMED], "« je prends … » sur le fil du forum, sans dépôt pour l'instant ; l'annonce expire au bout de 30 jours"],
        [STATUS_FR[FREE], "personne dessus à notre connaissance"],
        [],
        ["Correspondance", "Équivalente : couvre l'intégration HA. Partielle : en couvre une partie. "
                           "Automatique : rapprochée parce qu'elle porte le même nom."],
        ["Hors liste", "intégration Gladys sans entrée dans la liste : la liste a retiré ce que Gladys "
                       "couvrait déjà, et certaines n'existent pas chez HA"],
        [],
        ["Sources"],
        ["Liste Home Assistant", link(GIST_URL)],
        ["Store Gladys", link(STORE_PAGE)],
        ["Fil forum", link(FORUM_TOPIC)],
        ["Dépôts GitHub", link(f"https://github.com/topics/{TOPIC}", f"topic {TOPIC}")],
    ]:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14)
    for cell in ("A8", "A18"):
        ws[cell].font = Font(bold=True)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 110

    fills = {STATUS_FR[STORE]: "C6EFCE", STATUS_FR[TOPIC_ONLY]: "FFEB9C",
             STATUS_FR[NO_TOPIC]: "FCE4D6", STATUS_FR[INACTIVE]: "E7E6E6", STATUS_FR[CLAIMED]: "DDEBF7",
             STATUS_FR[NATIVE]: "D9E1F2"}
    sheets = [
        ("Liste HA ↔ Gladys", head_ha, rows_ha, 9, {8: "doc", 13: None}, 14,
         [6, 5, 34, 24, 9, 12, 9, 14, 6, 26, 13, 26, 16, 38, 22, 9, 10, 34, 60, 30]),
        ("Intégrations Gladys", head_g, rows_g, 4, {2: None}, 3,
         [28, 16, 40, 22, 26, 13, 9, 11, 11, 40, 30, 70]),
    ]
    for title, head, rows, status_col, link_cols, cover_col, widths in sheets:
        ws = wb.create_sheet(title)
        ws.append(head)
        for r in rows:
            r = list(r)
            for col, label in link_cols.items():
                if r[col]:
                    r[col] = link(r[col], label or r[col].replace("https://github.com/", ""))
            r[cover_col] = f'=IMAGE("{r[cover_col]}")' if r[cover_col] else ""
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for i, r in enumerate(rows, 2):
            if r[status_col] in fills:
                ws.cell(i, status_col + 1).fill = PatternFill("solid", fgColor=fills[r[status_col]])
            if r[cover_col]:
                ws.row_dimensions[i].height = 60
        ws.freeze_panes = "D2" if title.startswith("Liste") else "B2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    data = build()
    OUT.mkdir(exist_ok=True)
    (OUT / "data.json").write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    tables = flat_tables(data)
    for name, (head, rows) in zip(("liste-ha-gladys.csv", "integrations-gladys.csv"), tables):
        with open(OUT / name, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(head)
            w.writerows(rows)
    write_xlsx(data, tables, OUT / "suivi-integrations-gladys.xlsx")
    shutil.copy(ROOT / "index.html", OUT / "index.html")

    n_taken = sum(1 for h in data["ha"] if h["status"] != FREE and h["source"] == "list")
    july = [h for h in data["ha"] if h["source"] == "forum"]
    n_dup = sum(1 for g in data["gladys"] if g["dups"])
    unmapped = [g["repo"] for g in data["gladys"] if not g["ha"] and g["status"] != STORE]
    print(f"{len(data['ha']) - len(july)} HA ideas ({n_taken} taken), "
          f"{len(july)} July entries ({sum(1 for h in july if h['status'] != FREE)} covered), {len(data['gladys'])} Gladys repositories, "
          f"{n_dup} duplicated Gladys integrations; {len(unmapped)} repositories outside the store and not mapped")


if __name__ == "__main__":
    main()
